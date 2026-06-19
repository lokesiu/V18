"""Check quality of all generated documents."""
import sys, os, re
sys.path.insert(0, ".")
from docx import Document

customer_dir = r"D:\codex\V18\outputs\test_real_case_001\customer"

def check_docx(path):
    """Check a single DOCX file for quality issues."""
    issues = []
    try:
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        cn_count = sum(1 for ch in text if "\u4e00" <= ch <= "\u9fff")

        # Check Chinese char count
        if cn_count < 100:
            issues.append(f"Chinese chars too few: {cn_count}")

        # Check for placeholders
        for p in ["待补充", "TODO", "请自行补充", "XXX", "PLACEHOLDER"]:
            if p in text:
                issues.append(f"Placeholder found: {p}")

        # Check for markdown
        for p in ["**", "###"]:
            if p in text:
                issues.append(f"Markdown found: {p}")

        # Check for LLM conversational opening
        if text.startswith("好的") or "资深诉讼律师" in text[:100]:
            issues.append("LLM conversational opening detected")

        # Check for vague law citations
        vague = re.findall(r'《[^》]+》相关规定', text)
        if vague:
            issues.append(f"Vague law citations: {len(vague)}")

        return cn_count, issues, text[:500]
    except Exception as e:
        return 0, [f"Read error: {e}"], ""

print("=" * 60)
print("全量文书质量检查")
print("=" * 60)

for fname in sorted(os.listdir(customer_dir)):
    if not fname.endswith(".docx"):
        continue
    fpath = os.path.join(customer_dir, fname)
    cn_count, issues, preview = check_docx(fpath)
    status = "PASS" if not issues else "FAIL"
    print(f"\n[{status}] {fname}")
    print(f"  Chinese chars: {cn_count}")
    if issues:
        for i in issues:
            print(f"  [ISSUE] {i}")
    else:
        print("  No issues found")

print("\n" + "=" * 60)
print("检查完成")
print("=" * 60)
