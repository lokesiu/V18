"""Check quality of generated defense document."""
import sys
sys.path.insert(0, ".")
from docx import Document

doc_path = r"D:\codex\V18\outputs\test_real_case_001\customer\06_答辩状_上海嘉忞贸易有限公司案.docx"
doc = Document(doc_path)
full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

print("=" * 60)
print("答辩状质量检查")
print("=" * 60)

# Check 1: Contains defendant info
print("\n[1] 当事人信息:")
if "上海嘉忞贸易有限公司" in full_text:
    print("  [OK] 被告名称存在")
else:
    print("  [FAIL] 被告名称缺失")

if "李威璇" in full_text:
    print("  [OK] 原告名称存在")
else:
    print("  [FAIL] 原告名称缺失")

# Check 2: Law citations
print("\n[2] 法律引用:")
import re
law_citations = re.findall(r'《[^》]+》第[一二三四五六七八九十百千\d]+条', full_text)
for cite in law_citations:
    print(f"  - {cite}")
if not law_citations:
    print("  [WARN] 未找到具体法条引用")

# Check 3: No placeholder text
print("\n[3] 占位符检查:")
placeholders = ["待补充", "TODO", "请自行补充", "XXX", "PLACEHOLDER", "{{"]
for p in placeholders:
    if p in full_text:
        print(f"  [FAIL] 发现占位符: {p}")
    else:
        print(f"  [OK] 无 {p}")

# Check 4: Defense arguments
print("\n[4] 核心论点:")
arguments = [
    ("律师费抗辩", "律师费" in full_text and ("扩大" in full_text or "必要" in full_text)),
    ("金额争议", "3886" in full_text or "金额" in full_text),
    ("违约金", "违约金" in full_text),
    ("案件事实", "事实" in full_text),
    ("证据质证", "证据" in full_text),
]
for name, found in arguments:
    print(f"  [{'OK' if found else 'WARN'}] {name}")

# Check 5: Document structure
print("\n[5] 文档结构:")
sections = ["答辩请求", "事实与理由", "法律依据", "此致"]
for s in sections:
    if s in full_text:
        print(f"  [OK] {s}")
    else:
        print(f"  [WARN] 缺少 {s}")

# Check 6: Chinese character count
cn_count = sum(1 for ch in full_text if "\u4e00" <= ch <= "\u9fff")
print(f"\n[6] 中文字符数: {cn_count} ({'OK' if cn_count >= 100 else 'FAIL'})")

# Check 7: No markdown
print("\n[7] Markdown格式:")
md_patterns = ["**", "###", "- ", "* "]
for p in md_patterns:
    if p in full_text:
        print(f"  [WARN] 发现Markdown: {p}")
    else:
        print(f"  [OK] 无 {p}")

print(f"\n总字符数: {len(full_text)}")
print("=" * 60)
